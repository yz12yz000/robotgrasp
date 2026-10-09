from pathlib import Path
import json, hashlib, shutil
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]
ASSETS=HERE/'assets'
SOURCES={
'S1':'docs/rim_locator_第一版设计说明.md',
'S2':'docs/supported_rim_contacts_20260923.md',
'S3':'docs/demo_implementation_20260925.md',
'S4':'docs/live_demo_verification_20260927.md',
'S5':'src/rim_locator/rim_locator/rim_detection.py',
'S6':'src/rim_locator/config/place_config.json',
'S7':'debug_output/demo_implementation_20260925/verification_summary.json',
'S8':'debug_output/live_demo_20260927_gjgip1/live_verification_summary.json',
'S9':'debug_output/live_demo_20260927_gjgip1/execute_02_after_fix/result.json',
'S10':'docs/planning_diagnosis_20260923.md',
'S11':'debug_output/demo_implementation_20260925/moveit/result.json',
'S12':'src/grasp_executor/config/place_config.json',
}
IMAGES={
'I1':'debug_output/live_demo_20260927_gjgip1/initial_observation/rgb.png',
'I2':'debug_output/rim_geometry_20260923/frame01/geometry_before.png',
'I3':'debug_output/rim_geometry_20260923/frame01/contacts_detail.png',
'I4':'debug_output/live_demo_20260927_gjgip1/lift_02.png',
'I5':'debug_output/live_demo_20260927_gjgip1/final_scene.png',
'I6':'debug_output/rim_geometry_20260923/frame02/contacts_detail.png',
'I7':'debug_output/rim_geometry_20260923/frame03/contacts_detail.png',
'I8':'debug_output/live_demo_20260927_gjgip1/plan_01/round_0001/detection_overlay.png',
}
for key,src in IMAGES.items():
    shutil.copy2(ROOT/src,ASSETS/f'{key}.png')
manifest={'snapshot_date':'2026-10-02','evidence_dates':'2026-09-23 to 2026-09-27','generated_images':False,
          'note':'Original experiment images embedded unretouched. PowerPoint placement may crop the visible area. No new hardware tests were run for this presentation.',
          'sources':{},'images':{}}
for kind,items in [('sources',SOURCES),('images',IMAGES)]:
    for key,src in items.items():
        p=ROOT/src;manifest[kind][key]={'path':src,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
frames=[]
for i in range(1,4):
    path=f'debug_output/rim_geometry_20260923/frame0{i}/localization.json'
    data=json.loads((ROOT/path).read_text());frames.append(data)
    manifest['sources'][f'F{i}']={'path':path,'sha256':hashlib.sha256((ROOT/path).read_bytes()).hexdigest()}
(HERE/'evidence_sources.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))

prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
BG='F6F7F8'; INK='15283C'; MUTED='536678'; TEAL='087F83'; LIGHT='E8F2F2'; AMBER='A36416'; WHITE='FFFFFF'; LINE='D9E0E5'
FONT='Liberation Sans'
notes=[]

def shape(s,x,y,w,h,color,round=False):
    a=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    a.fill.solid();a.fill.fore_color.rgb=RGBColor.from_string(color);a.line.fill.background()
    if round:a.adjustments[0]=.12
    return a

def txt(s,text,x,y,w,h,size=20,color=INK,bold=False,align=PP_ALIGN.LEFT):
    box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=box.text_frame
    tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
    for j,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=line;p.alignment=align
        p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(5)
    return box

def pic(s,key,x,y,w,h,crop=False):
    path=ASSETS/f'{key}.png';iw,ih=Image.open(path).size
    if crop:
        p=s.shapes.add_picture(str(path),Inches(x),Inches(y),Inches(w),Inches(h))
        r=(w/h)/(iw/ih)
        if r<1:p.crop_left=p.crop_right=(1-r)/2
        else:p.crop_top=p.crop_bottom=(1-1/r)/2
    else:
        scale=min(w/iw,h/ih);pw,ph=iw*scale,ih*scale
        p=s.shapes.add_picture(str(path),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),width=Inches(pw),height=Inches(ph))
    return p

def base(num,section,title,subtitle=None):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
    shape(s,0,0,13.333,.12,TEAL)
    txt(s,section.upper(),.55,.31,12,.3,11,TEAL,True)
    txt(s,title,.55,.80,12.2,.66,30,INK,True)
    if subtitle:txt(s,subtitle,.55,1.48,12.15,.45,16,MUTED)
    shape(s,.55,7.03,12.23,.012,LINE)
    txt(s,'RIM_LOCATOR  /  RESEARCH PROGRESS  /  02 OCT 2026',.55,7.17,8,.18,9,MUTED)
    txt(s,f'{num:02d} / 06',11.85,7.15,.9,.22,10,TEAL,True,PP_ALIGN.RIGHT)
    return s

def cite(s,t):txt(s,t,.57,6.76,12.1,.2,9,MUTED)
def add_notes(s,title,body,refs):
    entry=f'第 {len(notes)+1} 页｜{title}\n\n{body}\n\n资料依据：\n'+ '\n'.join(f'{k}: {SOURCES.get(k,IMAGES.get(k,manifest["sources"].get(k,{}).get("path","")))}' for k in refs)
    s.notes_slide.notes_text_frame.text=entry
    notes.append({'slide':len(notes)+1,'title':title,'notes':body,'sources':refs})

def row(s,label,body,x,y,w=5.0):
    txt(s,label,x,y,w,.30,16,TEAL,True);txt(s,body,x,y+.38,w,.72,19)

# 1
s=base(1,'Scope & contribution','Rim-Guided Bowl Grasping','From a segmented point cloud to an executable grasp pose')
pic(s,'I1',5.55,2.08,7.23,4.40,crop=True)
shape(s,.55,2.09,4.62,4.39,WHITE,True)
row(s,'INPUT','Instance-labelled RGB-D points\n+ acquisition-time robot TF',.80,2.34,4.08)
row(s,'CORE CONTRIBUTION','A lightweight, geometry-based\nsymbolic model of rim contact',.80,3.63,4.08)
row(s,'OUTPUT','XYZ + orientation + diagnostics\nfor downstream MoveIt planning',.80,4.92,4.08)
cite(s,'Sources: S1, S3, S4, S5  |  Original lab image: I1 (27 Sep 2026)')
add_notes(s,'项目定位与贡献',
'这一页先说明 rim_locator 的边界：它是 ROS 2 中连接感知与运动规划的几何模块，不是 YOLO 模型，也不是机械臂驱动。上游负责分割和深度归属，它接收各实例的点云，并利用采集时刻的 TF 转入机器人基坐标系。输出除了位置，还包含四元数姿态和可解释的几何诊断。\n\n这里的 symbolic model 指基于碗沿、局部壁面、夹爪轴向和几何阈值的一组明确规则，不是语言模型或重新训练的抓取网络。右图是 9 月 27 日现场原始相机图，设备为 UR5e、双指夹爪和 RGB-D 相机，待抓物为绿色碗和浅色碗。今天汇报的重点是方法如何从单点定位扩展为可验证的抓放接口。',['S1','S3','S4','S5','I1'])

# 2
s=base(2,'Motivation & baseline','A detected bowl is not yet a supported grasp','The original height-band method was simple, but local wall evidence could fail.')
pic(s,'I2',.55,2.10,8.00,3.90)
txt(s,'Recorded single-patch diagnostic: rim points in red; wall support in orange.',.70,6.10,7.8,.50,13,MUTED)
shape(s,8.87,2.10,3.91,4.39,WHITE,True)
row(s,'ORIGINAL METHOD','98th-percentile height\n15 mm top band\nNear-side median → XYZ',9.12,2.32,3.36)
row(s,'OBSERVED FAILURE','Two masks detected;\nboth wall fits rejected\nin one recorded snapshot.',9.12,3.96,3.36)
txt(s,'Segmentation alone cannot\nvalidate a rim contact.',9.12,5.73,3.36,.64,17,AMBER,True)
cite(s,'Sources: S1, S2  |  Original point-cloud diagnostic: I2 (23 Sep 2026)')
add_notes(s,'原方法与问题证据',
'最初方法用 Z 的 98 分位数估计顶部，保留约 15 毫米高度带，在靠近基座的一侧选择局部点并取 XYZ 中位数，姿态由下游固定。对于较浅的碗，固定高度带容易包含碗壁，代表点可能偏低。\n\n图中展示的是随后加入单个壁面邻域拟合时保存的实际诊断图，而不是原始 XYZ 版本的成功率对比。一次记录中两个分割实例都正确，但绿色碗旧局部拟合倾角约 83.0°，浅色碗约 50.9°，都超过可接受范围。这个角度不等于真实碗口倾角，更可能反映邻域中混入过渡曲面或碗底。结论是不能只靠检测框和一个薄碗沿去声称接触方向可靠，需要给局部壁面提供证据。',['S1','S2','I2'])

# 3
s=base(3,'Geometry & pose design','Accept contacts only when the wall supports them','Same geometric rules for both bowls; no colour-specific model.')
pic(s,'I3',.62,2.10,3.18,4.34)
txt(s,'Projected contacts on a real capture',.62,6.48,3.30,.19,10,MUTED)
row(s,'01  ADAPT THE RIM BAND','b = clip(0.20 × H, 6, 12) mm\nH: robust height span (mm); upper-band anchor.',4.15,2.12,8.24)
row(s,'02  VERIFY THE LOCAL WALL','Fit at 12.5 / 18.75 / 25 mm radii.\nRequire agreement at two supported scales.',4.15,3.37,8.24)
txt(s,'≥20 points   •   ≥6 mm height span   •   ≤2.5 mm fit residual',4.15,4.52,8.2,.38,16,TEAL,True)
row(s,'03  BUILD THE TOOL POSE','Tool Y: closing direction. Tool Z: downward.\nQuaternion for control; RPY for inspection.',4.15,5.06,8.24)
shape(s,4.15,6.15,8.60,.42,LIGHT,True)
txt(s,'Current demo: tool tilt capped at 0°; measured wall tilt is logged separately.',4.31,6.22,8.25,.23,13,TEAL,True)
cite(s,'Sources: S2, S5, S6, S10  |  Original projected-contact image: I3')
add_notes(s,'几何规则与姿态构造',
'先在每个实例内部用连通性过滤去除明显离群点，再按稳健高度跨度自适应确定 6 到 12 毫米的半带宽。当前代码在第 98 分位高度上下各保留这个带宽，不是简单取最大高度。锚点取局部较高部分的中位数。\n\n原接触区域不足时，在真实观测到的顶部点带中进行有上限的搜索，间距 8 毫米，最多追加 96 个搜索区域。对每个区域尝试三种半径；至少两个尺度有新增点支持、法向相差不超过 12°才通过。还检查点数、纵向跨度、平面残差和锚点到平面的距离。法向缺失时拒绝，不编造竖直法向。\n\n姿态约定必须讲清：两指开合对应工具 Y 轴，工具 Z 向下，X 轴由右手系补齐。算法支持限制工具倾角，但当前演示配置 max_tool_tilt_deg=0，所以实机主要使用壁面水平投影确定夹合方向，工具保持竖直。图是早期采集的接触投影，只展示位置，不代表当前实机使用了对应的自由倾斜姿态。',['S2','S5','S6','S10','I3'])

# 4
s=base(4,'Implemented extensions','From a rim point to a complete pick-and-place task','Geometry stays in rim_locator; orchestration and motion remain separate.')
# parallel implementation changes, not a conceptual flow chart
rows=[('rim_locator','Up to 8 poses per bowl; 4 mm separation.\nVersioned groups + per-candidate diagnostics.'),
      ('ROS interfaces','Preserve instance IDs, source time and TF.\nKeep PointStamped / PoseArray compatibility.'),
      ('Execution & supervision','Try alternatives before any motion.\nRe-observe after release and return; stop on execution failure.')]
for i,(label,body) in enumerate(rows):
    y=2.12+i*1.21
    shape(s,.55,y,7.02,1.12,WHITE,True)
    txt(s,label,.76,y+.12,6.6,.28,16,TEAL,True)
    txt(s,body,.76,y+.47,6.57,.52,16)
pic(s,'I4',7.92,2.12,4.86,3.60,crop=True)
txt(s,'Original green-bowl grasp image • 27 Sep',7.96,5.81,4.8,.29,11,MUTED)
shape(s,.55,6.00,7.02,.58,LIGHT,True)
txt(s,'10 cm approach/lift   |   20 cm placement drop   |   1 s dwell',.74,6.17,6.65,.23,15,TEAL,True)
txt(s,'Fixed placement pose;\nnearest bowl first in each new observation.',7.95,6.13,4.7,.49,14,INK,True)
cite(s,'Sources: S3, S4, S6, S12  |  Original image: I4  |  Candidate fallback is pre-motion only.')
add_notes(s,'主要功能扩展与模块改动',
'这一页区分项目演进与早期对话方案：9 月 25 日版本已经从每碗单候选扩展到最多 8 个，按原几何搜索顺序保留，间距至少 4 毫米，没有新建综合评分器。rim_detection.py 负责候选；新增 candidate_protocol.py 校验分组协议；main.py 和 pointcloud.py 保留实例归属、原采集时间及诊断。原 PoseArray 仍只为每碗提供首选位姿，多候选通过 grasp_candidates 的版本化 JSON 传递。\n\n执行器只在尚未动作时允许换候选，整个六段路径必须先规划通过。实际执行每轮选择最近一只碗：接近、张开、下降、闭合、抬升、转移至固定准备位、沿基座 Z 下降 20 厘米、到位等待 1 秒、释放、抬升返回。当前接近与抬升高度是 10 厘米，不是早期的 15 厘米。\n\n完成后重新采集，再确定剩余物体；不把旧坐标或 YOLO ID 当跨帧身份。任务预算、命名空间、日志和运行器存活监督由 tools/test_live_place_pipeline.py 及执行器实现。桌面和额外夹爪保护盒已移除；原有机器人自碰撞与轨迹检查仍保留。',['S3','S4','S6','S12','I4'])

# 5
s=base(5,'Perception & software evidence','Real captures support the geometric design','Three recorded two-bowl observations; replay results are not a grasp success rate.')
# table built from original localization JSON, no fabricated metrics
x=.55;y=2.16;widths=[1.02,1.58,1.75,1.82,1.73];headers=['Capture','Accepted','Green tilt','Light tilt','Residual*']
for j,(w,hdr) in enumerate(zip(widths,headers)):
    xx=x+sum(widths[:j]);shape(s,xx,y,w,.49,INK);txt(s,hdr,xx+.12,y+.13,w-.2,.22,13,WHITE,True)
for i,frame in enumerate(frames):
    vals=frame['objects'];a=vals[0]['geometry'];b=vals[1]['geometry'];lo=min(a['wall_residual'],b['wall_residual'])*1000;hi=max(a['wall_residual'],b['wall_residual'])*1000
    data=[f'0{i+1}','2 / 2',f'{a["wall_tilt_deg"]:.2f}°',f'{b["wall_tilt_deg"]:.2f}°',f'{lo:.2f}–{hi:.2f} mm']
    yy=y+.49+i*.62
    for j,(w,v) in enumerate(zip(widths,data)):
        xx=x+sum(widths[:j]);shape(s,xx,yy,w,.62,WHITE if i%2==0 else 'EAF0F2');txt(s,v,xx+.12,yy+.18,w-.18,.25,15,INK,j==1)
txt(s,'* Local plane-fit residual, not independently measured localization accuracy.',.55,4.66,7.95,.49,13,MUTED)
for xx,big,label in [(0.55,'503','Regression tests passed\n27 Sep record'),(3.23,'15','MoveIt integration checks\n25 Sep; simulated control'),(5.91,'26','Historical replay checks\n25 Sep record')]:
    txt(s,big,xx,5.27,2.3,.57,33,TEAL,True);txt(s,label,xx,5.98,2.45,.53,12,MUTED)
for key,xx in [('I3',9.15),('I7',11.02)]:pic(s,key,xx,2.15,1.68,3.10)
txt(s,'Original contact projections\nCapture 01 / Capture 03',9.08,5.49,3.61,.69,15,MUTED)
cite(s,'Sources: F1–F3, S2, S7, S8, S11  |  Original images: I3, I7  |  All test counts are dated records.')
add_notes(s,'定位实验与软件验证',
'表格数字直接读取三份 localization.json，而不是从报告中手工估算。三次观测都为两个碗生成了有效位置，共六次定位接受。绿色碗估计法向倾角约 8.82 到 13.82 度，浅色碗约 31.89 到 37.37 度。它们是可见局部壁面的拟合量，不是独立标定的真实壁角，不能据此声称角度误差。\n\n六个局部拟合残差约 0.46 到 0.74 毫米，只说明所选点接近平面，不是系统定位精度，也不是抓取成功率。样本量仅三次观测，不报告泛化统计。右侧为同批次真实接触投影原图。\n\n下面三组数字来自不同日期和验证层级：9 月 27 日有 503 项回归通过，另 7 个可选启动测试单独通过；9 月 25 日有 15 项真实 MoveIt 配合模拟控制器检查，以及 26 项历史点云回放检查。它们不能相加为一次实验规模，也不等同于 503 次真实抓取。当前制作 PPT 没有重新启动机器人或重新跑硬件实验。',['F1','F2','F3','S2','S7','S8','S11','I3','I7'])

# 6
s=base(6,'Hardware evidence & next steps','One complete cycle; reliability remains to be measured','Documented case study: 27 September 2026')
for xx,key,label in [(0.55,'I1','Initial source scene'),(4.72,'I4','Green-bowl grasp'),(8.88,'I5','Final source scene')]:
    pic(s,key,xx,2.07,3.90,2.43)
    txt(s,label,xx,4.64,3.90,.34,16,INK,True)
shape(s,.55,5.15,6.08,1.38,WHITE,True)
txt(s,'AUDITED OUTCOME',.77,5.29,5.6,.26,13,TEAL,True)
txt(s,'Green: 1 complete cycle + fresh empty-ROI observation.\nWhite: lifted/placed; retreat cancelled → incomplete.',.77,5.73,5.56,.65,15)
shape(s,6.88,5.15,5.90,1.38,LIGHT,True)
txt(s,'NEXT VALIDATION',7.10,5.29,5.5,.26,13,TEAL,True)
txt(s,'Controlled repeats across bowls and poses.\nMeasure physical success; add grasp/release feedback.',7.10,5.73,5.38,.65,15)
cite(s,'Sources: S4, S8, S9  |  Original images: I1, I4, I5  |  No overall physical success rate claimed.')
add_notes(s,'实机结果与下一步',
'最后用现场原图和日志共同说明可行性。9 月 27 日的这组有完整记录的案例里，白碗确实被带起并放到左侧，但运行器异常结束，退回阶段被取消，因此不能计为完整流程成功。随后补上运行器身份监督与信号清理，再对剩余绿色碗采集新观测：生成 5 个候选，candidate 0 通过完整六段规划，实机完成抓取、转移、下降、释放和退回。下一轮新采集在取物 ROI 内检测到零目标，任务正常结束。\n\n这里只报告该案例修复后的一次完整成功，不把它说成整个项目只有一次动作，也不把白碗、其它手动会话和重复测试拼成成功率。目录中还存在控制器准备失败、过期反馈、取消和 remaining_unknown 等记录，场景与版本不统一，不能直接算总体成功率。空台面照片本身也不能单独证明抓稳和放稳。\n\n系统仍没有物体层面的传感反馈，日志明确保留 grasp_verified=false 和 placement_verified=false。现场可行候选都是 candidate 0，候选切换的硬件分支尚未在本案例中验证。下一步建议固定实验协议，跨碗型、位置和遮挡重复，分别统计检测接受、规划通过、实际抓起与稳定放置，并补充成功反馈。',['S4','S8','S9','I1','I4','I5'])

prs.core_properties.title='Rim-Guided Bowl Grasping — Research Progress'
prs.core_properties.subject='rim_locator: geometric contact support, pose generation, integration and experimental evidence'
prs.core_properties.author='Robot Grasp Project'
prs.core_properties.keywords='rim_locator, ROS 2, bowl grasping, geometry, Chinese speaker notes'
prs.save(HERE/'rim_locator_research_briefing_EN.pptx')
(HERE/'speaker_notes_zh.json').write_text(json.dumps(notes,ensure_ascii=False,indent=2))
(HERE/'speaker_notes_zh.txt').write_text('\n\n'+'\n\n'.join(f"第 {n['slide']} 页｜{n['title']}\n{n['notes']}\n资料：{', '.join(n['sources'])}" for n in notes))
(HERE/'analysis_and_structure.txt').write_text('''汇报设计：6 页，英文页面 + 中文演讲备注。
核心叙事：轻量几何规则将分割点云变成可解释、可规划的碗沿抓取位姿。
1. 项目定位与贡献：真实现场原图，说明输入、核心规则与输出。
2. 原方法与问题：区分早期 XYZ 高度带与后续单邻域壁面失败。
3. 当前几何方法：自适应带、多尺度证据、Y 轴闭合和工具倾角上限。
4. 功能扩展与模块边界：多候选、ROS 协议、有限重观测、抓放流程。
5. 定量实验：直接从三份原始 JSON 读取数据；软件测试按日期与层级区分。
6. 实机案例与限制：一次完整绿色碗流程、白碗退回取消，提出下一步实验。

关键核对：当前最多 8 候选；工具倾角限制 0 度；抬升 10 cm；放置下降 20 cm；到位停留 1 s。
图像策略：原图足够，全部使用原始实验图片。无 AI 生成图，无概念框架图，无未经标注的模拟图。
真实性边界：拟合残差不是定位误差；测试通过数不是抓取次数；无标准化总体物理成功率。
未修改机器人项目代码。evidence_sources.json 保存原始文件及 SHA256。
''')
print('Created',HERE/'rim_locator_research_briefing_EN.pptx')
