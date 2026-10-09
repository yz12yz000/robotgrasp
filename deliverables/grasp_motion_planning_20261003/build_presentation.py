from pathlib import Path
import json, hashlib, shutil, re
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
ASSETS = HERE / 'assets'
SOURCES = {
    'M1': 'src/grasp_executor/grasp_executor/moveit_ros.py',
    'M2': 'src/grasp_executor/grasp_executor/trajectory.py',
    'M3': 'src/grasp_executor/config/place_config.json',
    'M4': 'src/grasp_executor/grasp_executor/candidates.py',
    'D1': 'docs/demo_implementation_20260925.md',
    'D2': 'docs/live_demo_verification_20260927.md',
    'I1': 'debug_output/live_demo_20260927_gjgip1/initial_observation/rgb.png',
    'I2': 'debug_output/live_demo_20260927_gjgip1/lift_02.png',
}
manifest = {'date': '2026-10-03', 'sources': {}, 'generated_images': False,
            'note': 'Original laboratory photographs; no new hardware experiment. Notes contain explanations only.'}
for key, src in SOURCES.items():
    path = ROOT / src
    manifest['sources'][key] = {'path': src, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}
    if key.startswith('I'):
        shutil.copy2(path, ASSETS / (key + '.png'))
(HERE / 'evidence_sources.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2))

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
BG, INK, TEAL, MUTED = 'F6F7F8', '15283C', '087F83', '536678'
WHITE, LIGHT, LINE = 'FFFFFF', 'E8F2F2', 'D9E0E5'
FONT = 'Liberation Sans'
notes = []

def rect(s, x, y, w, h, color, rounded=False):
    a = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                           Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = RGBColor.from_string(color)
    a.line.fill.background()
    if rounded: a.adjustments[0] = .10
    return a

def text(s, t, x, y, w, h, size=20, color=INK, bold=False):
    a = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = a.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(t.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.font.name = FONT; p.font.size = Pt(size)
        p.font.bold = bold; p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(2)
    return a

def picture(s, key, x, y, w, h):
    path = ASSETS / (key + '.png')
    iw, ih = Image.open(path).size
    scale = min(w/iw, h/ih)
    pw, ph = iw*scale, ih*scale
    return s.shapes.add_picture(str(path), Inches(x+(w-pw)/2), Inches(y+(h-ph)/2),
                               width=Inches(pw), height=Inches(ph))

def slide(n, section, title, sub):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor.from_string(BG)
    rect(s, 0, 0, 13.333333, .12, TEAL)
    text(s, section.upper(), .55, .32, 12, .24, 11, TEAL, True)
    text(s, title, .55, .82, 12.3, .64, 31, INK, True)
    text(s, sub, .55, 1.52, 12.2, .47, 17, MUTED)
    rect(s, .55, 7.02, 12.23, .012, LINE)
    text(s, 'RIM_LOCATOR  /  MOTION PLANNING  /  03 OCT 2026', .55, 7.17, 10, .19, 9, MUTED)
    text(s, f'{n:02d} / 03', 12.05, 7.14, .8, .24, 11, TEAL, True)
    return s

def note(s, title, body):
    s.notes_slide.notes_text_frame.text = body
    notes.append({'slide': len(notes)+1, 'title': title, 'notes': body})

# Slide 1: pose, IK and the distinction between a target and a route.
s = slide(1, '01 / Inverse kinematics', 'From grasp pose to joint motion',
          'MoveIt 2 connects a geometric grasp target to an executable robot trajectory.')
items = [
    ('01  ALIGN THE TARGET', 'Convert grasp_center to tool0.\nUse measured joints as the starting state.'),
    ('02  SOLVE INVERSE KINEMATICS', 'Find joint angles that realize the tool pose.\nOne pose may have several arm configurations.'),
    ('03  PREFER NEARBY JOINT SOLUTIONS', 'Rank valid IK solutions by joint displacement.\nUp to 8 search attempts; plan up to 3 solutions.'),
]
for i, (label, body) in enumerate(items):
    y = 2.17 + i*1.27
    rect(s, .55, y, 6.65, 1.14, WHITE, True)
    text(s, label, .78, y+.13, 6.18, .27, 15, TEAL, True)
    text(s, body, .78, y+.49, 6.18, .56, 18)
picture(s, 'I1', 7.55, 2.17, 5.23, 3.27)
text(s, 'Original laboratory scene · 27 Sep 2026', 7.55, 5.55, 5.2, .26, 12, MUTED)
rect(s, 7.55, 5.98, 5.23, .62, LIGHT, True)
text(s, 'IK: joint angles. Planning: a route.', 7.75, 6.15, 4.84, .35, 16, TEAL, True)
text(s, 'FK(q) = T_target', .78, 6.14, 2.5, .34, 20, TEAL, True)
text(s, 'IK ranking:  J = Σᵢ(qᵢ − qᵢ,current)²', 3.35, 6.18, 3.85, .33, 16)
text(s, 'q: six joint angles    T: tool position + orientation    |    Sources: M1, M3; image: I1', .55, 6.76, 12, .19, 10, MUTED)
note(s, '从抓取位姿到关节运动',
'''抓取位姿说明的是夹爪要到哪里、朝向哪里，但是机械臂的电机需要知道六个关节分别转多少度。逆运动学，也就是 IK，就是把末端的位置和朝向换算成关节角度。人伸手拿杯子时，肘部可以抬高，也可以放低；同样，机械臂到达一个末端位姿，也可能有多种关节姿势。

程序首先把两指中间的夹持中心目标，换算成规划器使用的工具参考点目标。两者之间有固定偏移，转换之后才能让真正的夹持位置到达碗沿。同时读取机器人当前真实关节角度，确定从哪里开始运动。

公式 FK(q)=T 的意思是：六个关节角度决定工具的位置和朝向。IK 就是已知右边的目标，反过来求左边的角度。程序最多尝试八次不同的初始猜测，去掉重复或无效解，再按照各关节相对当前位置需要转动多少来排序。平方和越小，就越优先考虑；最多选三个解继续尝试规划。八次尝试并不等于一定得到八个不同解。

IK 只回答最后怎样摆放手臂。MoveIt 2 还要调用路径规划器，计算从当前姿势到目标姿势的中间运动。右侧照片是项目实际实验现场，用来对应机器人、夹爪和碗之间的关系。''')

# Slide 2: two algorithms, six concrete segments.
s = slide(2, '02 / Planning principles', 'Two planning methods, one six-segment task',
          'Joint-space search for transfer; Cartesian interpolation for vertical contact motions.')
cols = [(0.55, .48), (1.03, 3.85), (4.88, 2.25)]
rect(s, .55, 2.17, 6.58, .46, INK)
text(s, 'SEGMENT', .77, 2.30, 3.95, .23, 12, WHITE, True)
text(s, 'METHOD', 4.96, 2.30, 2, .23, 12, WHITE, True)
rows = [
    ('1', 'Approach: 10 cm above grasp', 'RRTConnect'),
    ('2', 'Descend to rim', 'Cartesian line'),
    ('3', 'Lift: +10 cm', 'Cartesian line'),
    ('4', 'Transfer to fixed ready pose', 'RRTConnect'),
    ('5', 'Lower: −20 cm', 'Cartesian line'),
    ('6', 'Return: +20 cm after release', 'Cartesian line'),
]
for i, row in enumerate(rows):
    y = 2.63+i*.50
    rect(s, .55, y, 6.58, .50, WHITE if i%2 == 0 else 'EDF1F4')
    text(s, row[0], .72, y+.13, .3, .25, 14, TEAL, True)
    text(s, row[1], 1.08, y+.13, 3.79, .30, 16)
    text(s, row[2], 4.97, y+.13, 2.07, .30, 15, TEAL, True)
rect(s, .55, 5.87, 6.58, .78, LIGHT, True)
text(s, 'Close after descent. At the drop pose:\nwait stationary for 1 s, then open and return.', .77, 6.02, 6.08, .56, 17, TEAL)
picture(s, 'I2', 7.55, 2.17, 5.23, 2.09)
text(s, 'Recorded green-bowl grasp · 27 Sep 2026', 7.55, 4.32, 5.23, .24, 11, MUTED)
text(s, 'RRTConnect / OMPL', 7.55, 4.79, 5.2, .29, 17, TEAL, True)
text(s, 'Sample and grow trees from start and goal;\nconnect valid states in six-joint space.', 7.55, 5.18, 5.2, .58, 17)
text(s, 'Cartesian interpolation', 7.55, 5.89, 5.2, .29, 17, TEAL, True)
text(s, 'Sample a tool-space line; solve IK along it.\n5 mm step; require 100% path completion.', 7.55, 6.26, 5.2, .56, 17)
text(s, 'Sources: M1, M3, D2; original image: I2', .55, 6.79, 6.6, .19, 10, MUTED)
note(s, '两种规划方法与六段动作',
'''完整抓放被拆成六段。第一段先移动到抓取点上方十厘米，接着张开夹爪，竖直下降到碗沿，闭合后再抬升十厘米。然后移动到事先保存的放置准备位，竖直下降二十厘米，到位并静止一秒后张开夹爪，最后抬升二十厘米返回准备位。

接近和搬运使用 RRTConnect。MoveIt 是组织规划和执行的软件框架，OMPL 是规划算法库，RRTConnect 是其中实际选择的算法。它可以理解为从出发点和终点两边探路，逐步扩展，直到两边接上。这里的一个点代表六个关节的一组角度，不是桌面上的一个位置；检查的是整条机械臂如何从一种姿势变到另一种姿势。

下降和抬升采用另一种方法：先规定夹持中心沿竖直线移动，再沿着这条线逐步求出关节角度。这叫笛卡尔路径计算。这样才能明确要求夹爪向下接近碗沿、向上提起，而不只是要求最终到达某个点。程序的路径插值步长为五毫米，这只是计算时的间隔，机器人不会每走五毫米就停一下。

只有整条直线路径都计算成功才接受。如果只算出九成路径，程序不会先执行九成再临时想办法。照片来自绿色碗实际抓取记录；六段动作采用的是同一套抓取与放置流程。''')

# Slide 3: actual validation and pre-motion selection.
s = slide(3, '03 / Validation and execution', 'Validate the full motion before executing it',
          'A candidate is executable only after all six segments pass the planning checks.')
rect(s, .55, 2.17, 8.06, .46, INK)
text(s, 'CHECK', .77, 2.30, 2.2, .22, 12, WHITE, True)
text(s, 'ACCEPTANCE RULE', 3.13, 2.30, 5.1, .22, 12, WHITE, True)
checks = [
    ('Reachability', 'Valid IK; target inside the workspace'),
    ('Sampled geometry', 'Joint bounds; workspace; scene collision checks'),
    ('Vertical motion', 'Complete path; bounded joint steps and lateral drift'),
    ('Endpoint accuracy', 'Position ≤ 2 mm; orientation ≤ 0.02 rad'),
    ('Speed and timing', 'Slow down to limits; reject if duration is excessive'),
    ('Before each segment', 'Start-state match; controller readiness; state recheck'),
]
for i, (label, rule) in enumerate(checks):
    y = 2.63+i*.51
    rect(s, .55, y, 8.06, .51, WHITE if i%2 == 0 else 'EDF1F4')
    text(s, label, .77, y+.14, 2.28, .29, 15, TEAL, True)
    text(s, rule, 3.13, y+.14, 5.26, .30, 15)
rect(s, 8.94, 2.17, 3.84, 3.54, WHITE, True)
text(s, 'PRE-MOTION SELECTION', 9.16, 2.38, 3.39, .30, 14, TEAL, True)
text(s, 'Try another IK solution\nfor the same grasp pose.', 9.16, 2.94, 3.36, .72, 20)
text(s, 'If full planning fails,\ntry the next grasp candidate.', 9.16, 3.89, 3.36, .76, 18)
text(s, 'Accept one complete plan\nbefore sending motion.', 9.16, 4.86, 3.36, .66, 17, TEAL, True)
rect(s, .55, 5.96, 12.23, .69, LIGHT, True)
text(s, 'EXECUTE', .77, 6.17, 1.3, .27, 15, TEAL, True)
text(s, 'Timed joint trajectory  →  MoveIt ExecuteTrajectory  →  Robot trajectory controller', 2.20, 6.17, 10.3, .32, 17)
text(s, 'Sources: M1–M4  |  Values shown are configured acceptance tolerances.', .55, 6.78, 12, .20, 10, MUTED)
note(s, '筛选路径与执行',
'''规划器返回一条路线以后，程序还要检查这条路线能不能用于当前任务。首先判断目标有没有有效关节解、是否在工作范围内。然后检查中间运动，而不只是看出发点和终点：关节不能超过允许角度，夹持中心不能越出工作范围，规划场景中的状态检查也必须通过。

对于下降和抬升，还要求路径完整、关节不能突然大幅跳变，夹爪不能明显向旁边偏移或改变朝向。到达终点时，计算结果与目标的位置偏差要求不超过两毫米，姿态偏差不超过零点零二弧度，约一点一五度。这些数值是程序接受方案的门槛，不是实验测得的定位精度。

如果路线可以走，但是速度或者加速度太大，程序会把整段运动放慢。放慢后如果耗时超过允许范围，也不能接受。对同一个抓取位姿，程序可以先换一种关节姿势重新规划；如果仍然无法形成完整可行方案，再尝试碗沿的下一个抓取候选。这些选择都发生在真实动作开始之前，而且六段运动要全部规划通过。

最终发送给机器人的是带时间的关节轨迹，也就是规定每个时刻六个关节应该在哪里、以多快的速度运动。MoveIt 把它交给机器人轨迹控制器执行。每段开始前还会确认机械臂实际位置与规划起点一致、控制器已准备好，并再次检查路径状态，避免把已经不适用的计划直接发给机器人。''')

prs.core_properties.title = 'Grasp Motion Planning: Principles and Validation'
prs.core_properties.subject = 'MoveIt 2, inverse kinematics, RRTConnect, Cartesian paths and trajectory validation'
prs.core_properties.author = 'Robot Grasp Project'
prs.save(HERE / 'grasp_motion_planning_EN.pptx')
(HERE / 'speaker_notes_zh.json').write_text(json.dumps(notes, ensure_ascii=False, indent=2))
(HERE / 'speaker_notes_zh.txt').write_text('\n\n'.join(f"第 {n['slide']} 页｜{n['title']}\n\n{n['notes']}" for n in notes))
assert len(prs.slides) == 3
for s in prs.slides:
    for sh in s.shapes:
        if sh.has_text_frame:
            assert not re.search(r'[\u4e00-\u9fff]', sh.text), sh.text
        assert sh.left >= 0 and sh.top >= 0
        assert sh.left + sh.width <= prs.slide_width + 100
        assert sh.top + sh.height <= prs.slide_height + 100
    assert re.search(r'[\u4e00-\u9fff]', s.notes_slide.notes_text_frame.text)
print('Created 3 English slides with explanatory Chinese notes.')
