from pathlib import Path
import json, re, hashlib
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from PIL import Image

H=Path(__file__).resolve().parent
ROOT=H.parents[1]
prs=Presentation(); prs.slide_width=Inches(13.333333); prs.slide_height=Inches(7.5)
INK='15283C'; TEAL='087F83'; WHITE='FFFFFF'; MUTED='536678'; LINE='D9E0E5'; LIGHT='E8F2F2'; BG='F6F7F8'
notes=[]

def box(s,x,y,w,h,c):
    a=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb=RGBColor.from_string(c); a.line.fill.background()
    # Remove template theme shadow for an academic, flat visual style.
    from pptx.oxml.xmlchemy import OxmlElement
    a._element.spPr.append(OxmlElement('a:effectLst'))
    for ref in a._element.xpath('./p:style/a:effectRef'):ref.set('idx','0')
    return a

def txt(s,t,x,y,w,h,size=18,c=INK,bold=False,center=False):
    a=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=a.text_frame
    tf.word_wrap=True; tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(t.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
        p.font.name='Liberation Sans'; p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(c)
        p.space_after=Pt(3)
        if center:p.alignment=PP_ALIGN.CENTER
    return a

def base(n,tag,title,sub):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
    box(s,0,0,13.333333,.10,TEAL)
    txt(s,tag,.55,.32,12,.25,11,TEAL,True)
    txt(s,title,.55,.80,12.3,.64,30,INK,True)
    txt(s,sub,.55,1.49,12.2,.45,16,MUTED)
    box(s,.55,7.02,12.23,.012,LINE)
    txt(s,'RIM_LOCATOR  /  MOTION PLANNING METHODS  /  03 OCT 2026',.55,7.16,11,.2,9,MUTED)
    txt(s,f'0{n} / 02',12.02,7.13,.8,.25,11,TEAL,True)
    return s

def pic_panel(s,left,right,x,y,w,h):
    path=H/'assets/planning_principles.png';iw,ih=Image.open(path).size
    ratio=iw*(1-left-right)/ih
    pw=min(w,h*ratio);ph=pw/ratio
    p=s.shapes.add_picture(str(path),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),Inches(pw),Inches(ph))
    p.crop_left=left;p.crop_right=right

def addnotes(s,title,body):
    s.notes_slide.notes_text_frame.text=body
    notes.append({'slide':len(notes)+1,'title':title,'notes':body})

s=base(1,'01 / PRINCIPLES','From a grasp pose to a planned motion',
       'MoveIt 2 combines inverse kinematics, joint-space search and Cartesian path computation.')
xs=[.55,4.74,8.93];w=3.85
for x in xs:box(s,x,2.13,w,4.42,WHITE)
for x,title,sub in zip(xs,['Inverse kinematics','RRTConnect / OMPL','Cartesian interpolation'],
                      ['POSE → JOINT CONFIGURATIONS','BIDIRECTIONAL SAMPLING','TOOL-SPACE WAYPOINTS']):
    txt(s,title,x+.17,2.31,w-.34,.35,21,TEAL,True)
    txt(s,sub,x+.17,2.78,w-.34,.22,10,MUTED,True)

box(s,.83,3.30,3.29,.67,LIGHT)
txt(s,'Target: position + orientation',.99,3.49,2.97,.25,15,TEAL,True,True)
txt(s,'FK(q) = T_target',.85,4.25,3.25,.48,25,INK,True,True)
txt(s,'q = [q₁, q₂, …, q₆]',.85,4.86,3.25,.34,19,MUTED,False,True)
txt(s,'One tool pose, multiple IK solutions.',.78,5.41,3.39,.56,16,INK,False,True)
pic_panel(s,0,.49,4.91,3.12,3.51,2.65)
pic_panel(s,.53,0,9.10,3.12,3.51,2.65)
txt(s,'Rank by J = Σᵢ(qᵢ − qᵢ,current)²\n8 IK attempts; plan up to 3 solutions.',.77,5.94,3.41,.56,14,TEAL)
txt(s,'Grow trees from start and goal.\nConnect valid joint states.',4.96,5.94,3.41,.56,16)
txt(s,'Line waypoints → successive IK.\n5 mm step; full path required.',9.15,5.94,3.41,.56,15)
txt(s,'Method illustrations: joint-space projection (centre), fixed-orientation tool path (right). Sources: M1, M3.',.55,6.78,12.23,.20,10,MUTED)
addnotes(s,'位姿如何变成规划路径',
'''夹爪的位姿说明它要到哪里、朝向哪里，但机械臂电机需要的是六个关节分别转多少度。左边的逆运动学，也叫 IK，就是把末端目标换算成关节角度。公式中，q 是六个关节角度，T 是工具的位置和朝向。人把手伸到同一个位置时，肘部可以高一些或低一些；机械臂也可能用不同关节姿势到达同一目标。程序先把夹持中心目标转换到规划使用的工具参考点，再从当前真实关节状态出发，最多尝试八次求解，按各关节需要转动的角度平方和排序，最多选三个解继续规划。

中间的图解释 RRTConnect。它从起点和终点两边分别扩展搜索树，尝试把两边连接起来。图中的每个点代表一组关节角度，深色折线表示连接成功后的一条路线。真实机械臂有六个关节，图中只用两个坐标轴帮助理解。MoveIt 2 是组织整个规划过程的软件框架，OMPL 是算法库，RRTConnect 是其中选择的搜索算法。IK 决定手臂最后怎么摆，路径搜索决定中间怎么走。

右边解释笛卡尔路径。程序先规定夹爪沿一条竖直线移动，再沿线逐步计算各个位置对应的关节角度。小坐标轴方向相同，表示这段升降保持工具朝向。五毫米是路径计算的插值步长，不是机器人每走五毫米停一下。只有整条路径都计算成功才接受。两张图都是方法原理示意，不是实际轨迹记录。''')

s=base(2,'02 / TASK PLANNING AND VALIDATION','Plan the complete task, then validate and execute',
       'Two joint-space transfers and four Cartesian segments are planned before any physical motion.')
stages=[('01','Approach','10 cm above rim','RRTConnect'),('02','Descend','To grasp pose','Cartesian'),
        ('03','Lift','+10 cm','Cartesian'),('04','Transfer','Fixed ready pose','RRTConnect'),
        ('05','Lower','−20 cm','Cartesian'),('06','Return','+20 cm','Cartesian')]
for i,(num,title,detail,method) in enumerate(stages):
    x=.55+i*2.06;bg=TEAL if method=='RRTConnect' else LIGHT;fg=WHITE if method=='RRTConnect' else INK
    box(s,x,2.10,1.93,1.22,bg)
    txt(s,num+'  '+title,x+.12,2.25,1.69,.30,17,fg,True)
    txt(s,detail,x+.12,2.68,1.69,.23,12,fg)
    txt(s,method,x+.12,3.02,1.69,.22,12,fg,True)
txt(s,'Gripper sequence: open before descent → close before lift → hold still for 1 s at drop → release → return.',.55,3.50,12.23,.40,15,MUTED)
cards=[
    ('01  STATE & GEOMETRY','Valid IK and workspace\nJoint bounds along the path\nSampled scene collision checks'),
    ('02  CARTESIAN FIDELITY','Complete path; no large joint jumps\nBounded lateral / angular deviation\nEndpoint: 2 mm and 0.02 rad'),
    ('03  TIMING & EXECUTION','Scale time to speed / acceleration limits\nReject excessive motion duration\nRecheck start state and controller'),
]
for i,(title,body) in enumerate(cards):
    x=.55+i*4.19;box(s,x,4.13,3.85,1.72,WHITE)
    box(s,x,4.13,3.85,.045,TEAL)
    txt(s,title,x+.16,4.34,3.53,.28,14,TEAL,True)
    txt(s,body,x+.16,4.82,3.53,.95,14)
txt(s,'PRE-MOTION SELECTION',.55,6.07,2.55,.28,12,TEAL,True)
txt(s,'Try another IK solution; if infeasible, try the next grasp candidate.',3.13,6.05,9.64,.33,17)
box(s,.55,6.51,12.23,.035,LINE)
txt(s,'Timed joint trajectory → MoveIt ExecuteTrajectory → Robot trajectory controller',.55,6.66,12.23,.30,17,TEAL,True)
addnotes(s,'完整抓放的规划、检查与执行',
'''一次抓放分为上面六段。先到碗沿上方十厘米，张开夹爪，向下到达抓取位置，闭合后抬升十厘米；然后移动到事先保存的放置准备位，竖直下降二十厘米，到位并静止一秒后释放，再向上返回准备位。深绿色两段采用 RRTConnect 搜索关节路径，浅绿色四段采用笛卡尔直线路径。程序在动作前先把六段全部规划出来，前一段的结束状态作为后一段的开始状态。

第一组检查关注能不能到达，以及中间状态是否有效。目标要有有效关节解，关节不能超过允许角度，夹持中心不能越出工作范围，还要通过规划场景中的碰撞检查。检查会覆盖中间轨迹的采样点，而不只是出发点和终点。

第二组检查关注升降是否符合要求。必须算出完整路径，关节不能突然大幅跳变，夹爪不能明显向旁边偏移或改变方向。路径终点的位置偏差要求不超过两毫米，朝向偏差不超过零点零二弧度，约一点一五度。这是程序接受规划结果的阈值，不是实测定位精度。

第三组检查关注速度和执行条件。路线可行但太快时，程序会延长运动时间来降低速度和加速度；如果减速后超过允许时长，就拒绝。一个抓取位姿不合适时，先尝试不同关节解；仍无法完成规划时，再换下一个抓取候选。换方案发生在真实运动之前。最终交给控制器的是带时间的关节轨迹，每段开始前还会检查实际关节状态与计划起点相符、控制器已准备好，并再次检查路径状态。''')

prs.core_properties.title='Grasp Motion Planning — Methods and Validation'
prs.core_properties.subject='Two-slide explanation of IK, RRTConnect, Cartesian planning and trajectory validation'
prs.core_properties.author='Robot Grasp Project'
prs.save(H/'grasp_motion_planning_EN_2slides.pptx')
(H/'speaker_notes_zh.json').write_text(json.dumps(notes,ensure_ascii=False,indent=2))
(H/'speaker_notes_zh.txt').write_text('\n\n'.join(f"第 {n['slide']} 页｜{n['title']}\n\n{n['notes']}" for n in notes))
sources={'M1':'src/grasp_executor/grasp_executor/moveit_ros.py','M2':'src/grasp_executor/grasp_executor/trajectory.py',
         'M3':'src/grasp_executor/config/place_config.json','M4':'src/grasp_executor/grasp_executor/candidates.py'}
manifest={'date':'2026-10-03','sources':{k:{'path':p,'sha256':hashlib.sha256((ROOT/p).read_bytes()).hexdigest()} for k,p in sources.items()},
          'figure':{'path':'assets/planning_principles.png','mode':'built-in image_gen','prompt':'image_prompt.txt',
                    'description':'Conceptual method illustration; not an experimental trajectory',
                    'sha256':hashlib.sha256((H/'assets/planning_principles.png').read_bytes()).hexdigest()},
          'robot_photographs':False}
(H/'evidence_sources.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
assert len(prs.slides)==2
for s in prs.slides:
    for sh in s.shapes:
        if sh.has_text_frame:assert not re.search(r'[\u4e00-\u9fff]',sh.text)
        assert sh.left>=0 and sh.top>=0 and sh.left+sh.width<=prs.slide_width+100 and sh.top+sh.height<=prs.slide_height+100
print('Created 2-slide English presentation with Chinese explanatory notes.')
