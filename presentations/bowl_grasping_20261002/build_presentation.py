"""Build an evidence-based progress deck. Does not start or modify ROS nodes."""
from pathlib import Path
import json
import re
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
NAME = 'UR5e_Bowl_Grasping_Progress_EN_CN'
LIVE = ROOT / 'debug_output/live_demo_20260927_gjgip1'
OLD = ROOT / 'debug_output/planning_diagnosis_20260923'
GEO = ROOT / 'debug_output/rim_geometry_20260923'
FONT = 'Liberation Sans'
C = dict(ink='142D43', muted='586A78', teal='007F82', green='447C36', blue='246FBC',
         light='F1F5F8', line='D7E0E7', white='FFFFFF', amber='A86A14', pale='F9F2E6')
P = Presentation()
P.slide_width, P.slide_height = Inches(13.333333), Inches(7.5)
P.core_properties.title = 'RGB-D–Guided Bowl Grasping and Placement with a UR5e'
P.core_properties.subject = 'Research progress; two-bowl perception, geometry, planning and supervised trials'
P.core_properties.author = 'Robot Grasping Project'
P.core_properties.keywords = 'UR5e, ROS2, MoveIt, bowl grasping, RGB-D, Chinese speaker notes'
RECORDS = []


def rgb(color):
    return RGBColor.from_string(C.get(color, color))


def rect(s, x, y, w, h, color, line=None):
    shape = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(color)
    shape._element.spPr.append(OxmlElement('a:effectLst'))
    for effect in shape._element.xpath('./p:style/a:effectRef'):
        effect.set('idx', '0')
    if line:
        shape.line.color.rgb = rgb(line); shape.line.width = Pt(.7)
    else:
        shape.line.fill.background()
    return shape


def text(s, content, x, y, w, h, size=20, color='ink', bold=False, align=None, font=FONT):
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(content.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = font; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = rgb(color)
        p.space_after = Pt(3)
        if align is not None: p.alignment = align
    return sh


def label(s, content, x, y, w, color='teal'):
    text(s, content.upper(), x, y, w, .23, 10, color, True)


def photo(s, path, x, y, w, h, view=None):
    path = Path(path)
    iw, ih = Image.open(path).size
    x0, y0, x1, y1 = view if view else (0., 0., float(iw), float(ih))
    # Display-only crop in PowerPoint: the embedded original is unchanged.
    vw, vh = x1-x0, y1-y0
    ratio = w/h
    if vw/vh > ratio:
        d = (vw-vh*ratio)/2; x0 += d; x1 -= d
    else:
        d = (vh-vw/ratio)/2; y0 += d; y1 -= d
    pic = s.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    pic.crop_left, pic.crop_right = x0/iw, 1-x1/iw
    pic.crop_top, pic.crop_bottom = y0/ih, 1-y1/ih
    RECORDS[-1]['images'].append(str(path.relative_to(ROOT)))
    return pic


def fit_picture(s, path, x, y, w, h):
    iw, ih = Image.open(path).size
    scale = min(w/iw, h/ih)
    return photo(s, path, x+(w-iw*scale)/2, y+(h-ih*scale)/2, iw*scale, ih*scale)


def card(s, title, body, x, y, w, h, accent='teal', title_size=22, body_size=17):
    rect(s, x, y, w, h, 'light')
    rect(s, x, y, .045, h, accent)
    text(s, title, x+.18, y+.14, w-.36, .42, title_size, accent, True)
    text(s, body, x+.18, y+.72, w-.36, h-.8, body_size)


def take(s, content, color='teal'):
    rect(s, .55, 6.57, 12.23, .39, 'light')
    text(s, content, .73, 6.64, 11.9, .25, 14, color, True)


def table(s, headers, rows, x, y, widths, row_h=.59, header_h=.5, font_size=17):
    xx=x
    for value,w in zip(headers,widths):
        rect(s,xx,y,w,header_h,'ink'); text(s,value,xx+.12,y+.12,w-.24,header_h-.12,12 if w<1.5 else 14,'white',True);xx+=w
    yy=y+header_h
    for j,row in enumerate(rows):
        xx=x
        for value,w in zip(row,widths):
            rect(s,xx,yy,w,row_h,'light' if j%2==0 else 'white')
            text(s,str(value),xx+.12,yy+.13,w-.24,row_h-.17,font_size)
            xx+=w
        yy+=row_h


def slide(title, section, subtitle='', sources=(), note='', tag='PROJECT EVIDENCE'):
    s=P.slides.add_slide(P.slide_layouts[6]); s.background.fill.solid();s.background.fill.fore_color.rgb=rgb('white')
    RECORDS.append(dict(number=len(P.slides),title=title,section=section,images=[],sources=list(sources),note=note))
    label(s,section,.55,.3,8)
    text(s,tag,9.1,.3,3.68,.22,10,'muted',align=PP_ALIGN.RIGHT)
    text(s,title,.55,.77,12.2,.66,30,'ink',True)
    if subtitle: text(s,subtitle,.57,1.43,12.1,.34,15,'muted')
    rect(s,.55,7.12,12.23,.012,'line')
    text(s,'UR5e  /  RGB-D BOWL MANIPULATION',.55,7.22,6,.16,9,'muted')
    text(s,f'{len(P.slides):02d}  /  18',11.78,7.2,1,.2,10,'muted',align=PP_ALIGN.RIGHT)
    return s


initial=LIVE/'initial_observation/rgb.png'
overlay=LIVE/'plan_01/round_0001/detection_overlay.png'
plan=json.loads((LIVE/'plan_01/round_0001/result.json').read_text())
batch=json.loads((LIVE/'plan_01/round_0001/candidates.json').read_text())
groups={o['object_id']:o for o in batch['objects']}
vision={o['object_id']:o for o in plan['statuses']['yolo_vision']['objects']}

# 01
s=slide('','RESEARCH PROGRESS  /  02 OCT 2026',sources=['docs/README.md','docs/live_demo_verification_20260927.md'],tag='EXPERIMENTS: 23–27 SEP 2026',note='本项目利用 RGB-D 相机识别桌面上的碗，根据碗沿局部几何生成抓取姿态，再由 UR5e 完成抓取与固定位置放置。今天重点汇报从单点抓取发展到同碗多候选和逐次重新观察的过程，并用白碗、绿碗的原始结果说明实际进展。所有实验图片来自项目记录，离线测试和现场运动会明确区分。本报告不是总体成功率统计。')
text(s,'RGB-D–Guided\nBowl Grasping\nand Placement',.6,1.28,5.1,2.3,37,'ink',True)
text(s,'with a UR5e',.64,3.66,4.8,.6,29,'teal',True)
text(s,'From single contacts to\ncandidate-based planning\nand repeated observation',.64,4.49,4.8,1.27,20,'muted')
photo(s,initial,5.93,1.08,6.85,5.46,view=(70,120,970,800))
label(s,'Two bowls · One robot · Verified execution',.65,6.22,5.1)

# 02
s=slide('A defined tabletop task, with separate pick and place regions','01  /  TASK & SETUP','UR5e + parallel gripper + Orbbec Gemini2L + ROS 2 Humble',sources=['docs/live_demo_verification_20260927.md','src/yolo_vision/config/live_pick_roi_20260927.json'],note='任务是在固定相机与机器人基座条件下，从指定取物区域内选择碗，夹持碗沿，并搬到保存好的放置位姿。图中的取物框对应 9 月 27 日现场 ROI，放置标记表示固定参考位置的近似图像投影，不表示已测量的容器边界。两区必须分开，否则下一轮可能重新识别刚放下的碗。当前系统不加载桌面或额外夹爪保护盒，仍保留机器人 URDF 自碰撞和轨迹约束。')
photo(s,initial,.55,1.95,7.3,4.5625)
# Native annotations in original image pixel coordinates.
sx=7.3/1280;sy=4.5625/800
roi=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(.55+610*sx),Inches(1.95+385*sy),Inches(265*sx),Inches(170*sy))
roi.fill.background();roi.line.color.rgb=rgb('teal');roi.line.width=Pt(2)
rect(s,.55+610*sx,1.95+352*sy,1.8,.24,'teal');text(s,'PICKING REGION',.58+610*sx,1.96+352*sy,1.75,.2,10,'white',True)
dot=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(.55+460*sx-.06),Inches(1.95+433*sy-.06),Inches(.12),Inches(.12));dot.fill.solid();dot.fill.fore_color.rgb=rgb('blue');dot.line.fill.background()
text(s,'Fixed placement\nreference',2.5,4.62,1.75,.55,13,'white',True)
card(s,'1  Detect','Observe the bowls inside the picking region.',8.13,1.95,4.65,1.34)
card(s,'2  Grasp and place','Use a rim contact and a stored placement pose.',8.13,3.48,4.65,1.34)
card(s,'3  Observe again','Acquire a new scene after release and retreat.',8.13,5.0,4.65,1.34)

# 03
s=slide('What changed from the initial grasping pipeline','02  /  CONTRIBUTIONS','Three ROS packages retained; perception, geometry and execution extended.',sources=['docs/README.md','docs/demo_implementation_20260925.md','docs/planning_diagnosis_20260923.md'],note='系统继续使用 yolo_vision、rim_locator 和 grasp_executor 三个包。主要改动是从单目标点云扩展到逐实例数据，从一个首选接触点扩展到多个有几何依据的候选，并把抓取扩展为完整六段抓放。运行器负责每次只执行一只碗，完成后重新观察。除此之外，还增加了配置快照、任务隔离、进程监督和可重复测试。使用的是已有分割模型，没有声称训练了新的抓取神经网络。')
table(s,['Area','Earlier capability','Current implementation'],[
['Perception','Single-target output','Multiple instances + depth diagnostics'],
['Rim geometry','One preferred contact','Up to 8 supported candidates per bowl'],
['Motion','Approach / descend / lift','Six-stage grasp, transfer and placement'],
['Task execution','Reuse one target snapshot','One bowl per round; then observe again'],
['Operation','Separate manual processes','Mode isolation, logs and process supervision']],.55,2.0,[2.03,4.0,6.2],row_h=.77,font_size=18)
take(s,'Contribution: an integrated, testable manipulation pipeline with explicit failure handling.')

# 04
s=slide('Two bowls are segmented and reconstructed independently','03  /  TWO-BOWL PERCEPTION','One synchronized RGB-D observation  |  27 Sep 2026  |  plan-only dataset',sources=['debug_output/live_demo_20260927_gjgip1/plan_01/round_0001/result.json'],note='这是同一帧真实双碗观测。白碗和绿碗共享类别 bowl，但拥有不同的观测内编号、分割掩膜和三维点集合。白碗有少量无效深度，绿碗在该帧掩膜内深度全部有效。程序保留原始采集时间，并避免掩膜重叠造成同一个像素分给两个物体。这里的置信度与有效深度比例只代表本帧，不代表整个系统的检测率。颜色名称仅用于向观众说明，不是算法分支。')
photo(s,overlay,.55,2.02,7.2,3.45,view=(560,350,900,570))
label(s,'Original segmentation overlay · display crop',.56,5.64,7.1)
table(s,['Recorded value','White bowl','Green bowl'],[
['Instance ID','1','0'],
['Confidence',f"{vision[1]['confidence']:.3f}",f"{vision[0]['confidence']:.3f}"],
['Valid depth points',f"{vision[1]['point_count']:,}",f"{vision[0]['point_count']:,}"],
['Valid depth ratio',f"{vision[1]['valid_depth_ratio']:.1%}",f"{vision[0]['valid_depth_ratio']:.1%}"]],8.04,2.02,[2.19,1.27,1.27],row_h=.7,font_size=16)
take(s,'Detection, valid depth and a usable grasp are three different checks.')

# 05
s=slide('White and green bowls offer different wall evidence','04  /  TWO-BOWL GEOMETRY','Recorded geometry example  |  23 Sep 2026, frame 01',sources=['debug_output/rim_geometry_20260923/frame01/localization.json','docs/supported_rim_contacts_20260923.md'],note='双碗试验的价值不仅是目标数量增加，也在于局部壁面和深度分布不同。这里展示的是 9 月 23 日的一次实际观测，绿色碗选中局部壁面的倾角约 11.7 度，白色碗约 31.9 度。这些角度是该次接触区域的拟合结果，不是每只碗固定不变的材料或形状参数。当前正式配置将工具倾角限制为零，即工具仍竖直接近，壁面几何提供接触位置及水平夹合方向。')
photo(s,GEO/'frame01/contacts.png',.55,1.97,5.84,3.13,view=(665,450,825,550))
photo(s,GEO/'frame01/contacts.png',6.64,1.97,6.14,3.13,view=(680,377,800,455))
card(s,'Green bowl','30 wall points  |  11.73° local wall tilt\n0.609 mm plane residual',.55,5.18,5.84,1.33,'green',20,16)
card(s,'White bowl','33 wall points  |  31.89° local wall tilt\n0.549 mm plane residual',6.64,5.18,6.14,1.33,'blue',20,16)
take(s,'Measured wall tilt is not commanded tool tilt; the current tool remains vertical.')

# 06
s=slide('Both bowls were localized in three recorded observations','05  /  TWO-BOWL LOCALIZATION','Recorded RGB-D replay  |  same geometric acceptance rules',sources=['docs/supported_rim_contacts_20260923.md','debug_output/rim_geometry_20260923/frame01/localization.json','debug_output/rim_geometry_20260923/frame02/localization.json','debug_output/rim_geometry_20260923/frame03/localization.json'],note='旧方法只在一个近侧锚点和固定邻域拟合壁面，容易受遮挡、弧面或碗底污染影响。改进后在顶部点带搜索其他接触区域，并要求不同尺度下的法向一致。三份独立现场观测都检测到两只碗，也都为两只生成了定位结果。表格展示局部壁面支持点数和倾角。六次拟合的残差约为 0.46 到 0.74 毫米，这只是拟合质量，不能作为绝对定位精度或实际抓取成功率。')
photo(s,GEO/'frame01/contacts.png',.55,1.99,4.3,3.75,view=(610,340,850,575))
label(s,'Measured contacts, frame 01',.56,5.9,4.4)
rows=[]
for f in ('frame01','frame02','frame03'):
    objects=json.loads((GEO/f/'localization.json').read_text())['objects'];g,w=[o['geometry'] for o in objects]
    rows.append([f[-2:],f"{g['wall_count']} / {g['wall_tilt_deg']:.2f}°",f"{w['wall_count']} / {w['wall_tilt_deg']:.2f}°",'2 / 2'])
table(s,['Frame','Green: points / tilt','White: points / tilt','Localized'],rows,5.16,2.04,[.83,2.42,2.42,1.95],row_h=.72,font_size=17)
text(s,'Quality gates',5.28,5.05,4.8,.4,21,'teal',True)
text(s,'≥20 wall points  ·  ≥6 mm wall span\n≤2.5 mm residual  ·  ≤12° normal change',5.28,5.59,7.1,.68,18)
take(s,'Observed fitting residual: 0.46–0.74 mm. This is not ground-truth positioning accuracy.')

# 07
s=slide('Each bowl has its own candidates and processing order','06  /  TWO-BOWL CANDIDATES','27 Sep plan-only snapshot  |  complete six-stage plans found for both bowls',sources=['debug_output/live_demo_20260927_gjgip1/plan_01/round_0001/candidates.json','debug_output/live_demo_20260927_gjgip1/plan_01/result.json'],note='这一帧白碗和绿碗各有 8 个抓取候选。候选属于同一只碗的不同接触选择，不会被当成额外目标。物体按首选接触点到机器人基座的 XY 距离排序，因此该帧白碗先于绿碗。这里的编号只属于当前观测，不能拿其他帧的编号直接匹配。两只碗各自完整六段规划都已通过，但这是 plan-only；真正执行时会重新观测，并且每轮只抓一只。')
photo(s,overlay,.55,1.95,5.64,2.77,view=(620,370,865,525))
text(s,'Ordering uses geometry, not color or instance ID.',.64,5.0,5.43,.85,23,'ink',True)
rows=[]
for name,oid in [('White bowl',1),('Green bowl',0)]:
    o=groups[oid];c=o['candidates'][0]
    rows.append([name,str(oid),str(len(o['candidates'])),f"{o['distance']:.3f} m",'6 / 6'])
table(s,['Bowl','ID','Candidates','XY distance','Plan'],rows,6.45,1.97,[1.68,.5,1.3,1.57,1.28],row_h=.86,font_size=17)
card(s,'Before the first motion','Reject infeasible candidates. Execute only after the entire six-stage path passes.',6.45,4.58,6.33,1.55,title_size=22,body_size=18)
take(s,'Both targets planned successfully in this snapshot; planning success is not physical grasp success.')

# 08
s=slide('Feasible contacts still require feasible robot motion','07  /  MOTION PLANNING','Bounded IK search + OMPL transfer + Cartesian vertical segments',sources=['docs/planning_diagnosis_20260923.md','src/grasp_executor/grasp_executor/ik_angles.py','debug_output/demo_implementation_20260925/moveit/result.json'],note='接触点几何合理之后，机械臂仍可能无法到达或无法完成后续下降。现场曾复现肩部和腕部关节绕行、软硬限位不一致及接近轨迹过长的问题。当前实现检查等价关节角、软硬限位交集，并对有限个 IK 分支按关节行程排序。接近、下降和抬升还要联合检查。同碗候选回退已有真实 MoveIt 配合模拟控制器的证据，但 9 月 27 日现场运行使用的是首候选，没有实机验证首候选失败后换第二个的分支。')
photo(s,LIVE/'lift_02.png',.55,1.98,6.95,4.34,view=(240,255,1000,730))
card(s,'Joint-space consistency','Equivalent angles and URDF soft limits prevent avoidable joint winding.',7.78,1.98,5.0,1.33,body_size=17)
card(s,'Whole-path feasibility','Validate approach, descent, lift, transfer, lowering and retreat.',7.78,3.5,5.0,1.33,body_size=17)
card(s,'Bounded fallback','Try another candidate only before motion has started.',7.78,5.02,5.0,1.33,body_size=17)
take(s,'Candidate fallback is verified in isolated MoveIt tests; live fallback remains untested.')

# 09
s=slide('A stored TCP reference makes placement repeatable','08  /  SIX-STAGE GRASP & PLACE','Fixed 6D pose  |  static TCP-to-grasp-center transform  |  MoveIt execution',sources=['src/grasp_executor/config/place_config.json','src/grasp_executor/grasp_executor/grasp.py','docs/place_offline_debug_20260923.md'],note='固定放置准备姿态保存于配置，以机器人固定基座为参考，不在每次程序启动时重新采当前位置。执行前将 TCP 参考通过静态 TF 转为 grasp_center 的目标，再由 MoveIt 求解。运动一共六段；接近位打开夹爪，下降后闭合，抬升后转移，放置下降到位并静止后等待一秒，再打开并退回。预抓和抬升高度为 10 厘米，放置下降为 20 厘米。页面照片是不同实验中的动作示例，不是同一次连续过程。')
for path,cap,x in [(OLD/'green_lifted.png','LIFT · 23 SEP',.55),(OLD/'green_placed.jpg','PLACE · 23 SEP',4.71),(OLD/'white_placed_final.png','RETREAT · 23 SEP',8.86)]:
    photo(s,path,x,2.0,3.91,2.82,view=(275,230,915,690));label(s,cap,x,4.94,3.9)
text(s,'01  Approach     02  Descend     03  Lift     04  Transfer     05  Lower     06  Retreat',.64,5.52,12.0,.42,18,'ink',True)
take(s,'10 cm approach / lift     |     20 cm placement descent     |     1 s hold before release')

# 10
s=slide('One completed bowl, then a new observation','09  /  REPEATED OBSERVATION','Post-fix green-bowl run  |  27 Sep 2026',sources=['debug_output/live_demo_20260927_gjgip1/execute_02_after_fix/result.json','debug_output/live_demo_20260927_gjgip1/live_verification_summary.json'],note='修复后的绿碗运行第一轮检测到一只碗，完成抓取、放置、释放和退回后，自动启动第二轮。第二轮原始采集时间严格晚于新的采集门槛，检测到取物区为空后任务结束。这里不复用旧目标坐标，也不依赖跨帧物体编号。右侧照片记录最终场景，零检测结论同时来自第二轮结果日志。它不能独立证明放置成功，必须结合执行状态和现场观察。这是分步重新观察，不是运动中视觉伺服。')
photo(s,LIVE/'execute_02_after_fix/round_0001/detection_overlay.png',.55,2.0,5.98,3.3,view=(345,330,900,650))
photo(s,LIVE/'final_scene.png',6.8,2.0,5.98,3.3,view=(345,330,900,650))
label(s,'Round 1 · target detected',.56,5.5,5.8)
label(s,'Round 2 · empty picking region',6.81,5.5,5.8)
text(s,'1 bowl completed\n6 motion stages executed',.56,5.88,5.8,.56,18)
text(s,'Fresh capture after retreat\n0 targets reported by the next observation',6.81,5.88,5.8,.56,18)
take(s,'New timestamps, new candidates and a measured robot start state for every round.')

# 11
s=slide('Both bowl types were lifted and placed in supervised trials','10  /  TWO-BOWL PHYSICAL EVIDENCE','23 Sep 2026  |  observations from separate execution and recovery steps',sources=['docs/planning_diagnosis_20260923.md'],note='这四张原图记录 9 月 23 日绿碗抓起、绿碗放下、白碗抓起和白碗放下。绿碗首次执行中有人进入区域而中止，之后经用户确认使用一次性恢复脚本继续完成，不能当作自动恢复能力。随后人工清空固定放置区，白碗正式执行最终完整成功。两类碗都完成实际搬运，有图像和日志支持；但是这些不是一次无人干预的连续双碗任务，也不能由这两次展示计算总体成功率。')
panels=[(OLD/'green_lifted.png','GREEN · LIFTED',.55,1.94),(OLD/'green_placed.jpg','GREEN · PLACED',6.8,1.94),(OLD/'white_lifted.jpg','WHITE · LIFTED',.55,4.2),(OLD/'white_placed_final.png','WHITE · PLACED',6.8,4.2)]
for path,cap,x,y in panels:
    photo(s,path,x,y,5.98,1.87,view=(240,260,940,630));label(s,cap,x,y+1.94,5.9,'green' if 'GREEN' in cap else 'blue')
take(s,'Human intervention occurred between these steps; this is not an uninterrupted two-bowl cycle.','amber')

# 12
s=slide('The fixed placement point exposed a multi-bowl limitation','11  /  TWO-BOWL TASK CONSTRAINT','A successful first placement can occupy the destination needed by the next bowl.',sources=['docs/planning_diagnosis_20260923.md','docs/demo_implementation_20260925.md'],note='双碗任务揭示了单碗演示不容易发现的问题：第一个碗放到固定位置后，第二个碗不能直接复用同一个被占据的放置空间。9 月 23 日的实验由人员移走已放好的绿碗，之后再继续白碗。现有程序没有放置区容量感知，也没有自动分配多个落点或堆叠策略。当前取物 ROI 与放置区分离，可以减少重复检测，但它本身并不解决接收区占用问题。这是后续真正连续多碗任务必须验收的条件。')
photo(s,OLD/'green_placed.jpg',.55,1.98,7.18,4.49,view=(170,180,1000,735))
card(s,'What happened','The placed green bowl occupied the fixed destination.',8.02,1.99,4.76,1.44,accent='amber',body_size=18)
card(s,'What the trial required','An operator cleared the destination before the white-bowl trial.',8.02,3.63,4.76,1.44,accent='amber',body_size=18)
text(s,'Not yet implemented:\nplacement occupancy sensing\nor automatic slot allocation',8.18,5.32,4.44,.93,18,'muted')
take(s,'Separating the picking ROI prevents re-picking; it does not prove placement capacity.','amber')

# 13
s=slide('Post-fix execution completed a full cycle and re-observation','12  /  LIVE VALIDATION  ·  27 SEP','Selected run: execute_02_after_fix  |  real robot and camera',sources=['docs/live_demo_verification_20260927.md','debug_output/live_demo_20260927_gjgip1/live_verification_summary.json','debug_output/live_demo_20260927_gjgip1/execute_02_after_fix/result.json'],note='9 月 27 日初始双碗都通过了六段规划。白碗现场抓起并放下，但首次运行终端异常退出，退回阶段被取消，所以不能算完整任务成功。修复运行器监督后，重新采集剩余绿碗并完成全流程，第二轮重新观察到零目标。这里表中的完成次数仅指这次修复后的绿碗运行，不是整个项目的实验总数。动作结果成功仍不等于夹持传感器确认，程序保留 grasp_verified=false 和 placement_verified=false。')
photo(s,LIVE/'lift_02.png',.55,1.98,6.85,3.8,view=(265,245,990,685))
label(s,'Green-bowl manipulation · original camera image',.56,5.97,6.8)
table(s,['Recorded field','Result'],[['Outcome','completed'],['Completed cycles','1'],['Observation rounds','2'],['Candidate attempts','1'],['Cleanup errors','None reported']],7.7,1.99,[2.71,2.37],row_h=.62,font_size=18)
take(s,'Command completion is logged; independent grasp / placement sensing is not available.')

# 14
s=slide('An interrupted trial led to stronger process supervision','13  /  FAILURE-DRIVEN IMPROVEMENT','The first 27 Sep white-bowl run was not counted as a complete success.',sources=['docs/live_demo_verification_20260927.md','src/grasp_executor/grasp_executor/process_guard.py'],note='第一次白碗运行的终端会话结束后，独立执行器仍在运行。终端被谁结束尚未确定，已经确认并修复的是执行器没有监测运行器消失这一缺口。新增监督机制同时检查 PID 和进程启动身份，防止 PID 复用；运行器消失时触发既有取消和停止流程，不自动释放。SIGTERM 与 SIGHUP 的所属节点清理已在隔离 ROS 中验证，运动期间异常由假接口测试，没有故意对运动中的真机注入崩溃。')
photo(s,LIVE/'lift_01.png',.55,1.98,5.5,3.44,view=(260,260,960,690))
label(s,'White bowl lifted · later retreat was cancelled',.56,5.59,5.6,'amber')
text(s,'Observed gap',6.39,2.05,5.9,.4,23,'amber',True)
text(s,'The runner exited while the independent executor remained active.',6.39,2.57,6.0,.85,21)
text(s,'Implemented response',6.39,3.68,5.9,.4,23,'teal',True)
text(s,'Process-identity monitoring\nExplicit termination-signal handling\nTask-owned process cleanup',6.39,4.24,6.0,1.3,21)
take(s,'After motion starts: stop and report errors; do not automatically release, switch candidates or regrasp.')

# 15
s=slide('Software validation is distinct from physical success','14  /  VALIDATION EVIDENCE','Recorded test results  |  suite counts are not combined into a grasp success rate',sources=['debug_output/live_demo_20260927_gjgip1/live_verification_summary.json','debug_output/demo_implementation_20260925/verification_summary.json'],note='验证分成纯逻辑、ROS 启动和通信、实际 MoveIt 配合模拟控制器，以及现场运动几个层级。503 项是 9 月 27 日软件回归数量，默认跳过的 7 项隔离启动测试后来单独全部通过。15 项 MoveIt 检查来自 9 月 25 日，包含候选回退、模拟执行、不可达拒绝和工具变换。表中日期及范围分别列出，不把不同套件数量相加，也不把软件通过数当成物理抓取次数。现场总体成功率仍需规范重复试验。')
table(s,['Validation layer','Result','Date','What it verifies'],[
['Regression tests','503 passed','27 Sep','Logic, state transitions and failure handling'],
['Isolated ROS startup','7 passed','27 Sep','Controller checks and startup behavior'],
['Launch contracts','32 checks','27 Sep','Parameters, namespaces and supervision'],
['MoveIt integration','15 checks','25 Sep','Real planning with simulated execution'],
['Physical trials','Images + logs','23 / 27 Sep','Supervised lifting, placement and retreat']],.55,2.0,[2.73,1.65,1.1,6.75],row_h=.76,font_size=17)
take(s,'Physical success rate remains unmeasured; automated test coverage is a separate result.')

# 16
s=slide('What is demonstrated, and what needs the next experiment','15  /  CONCLUSIONS & NEXT STEPS','A working manipulation system with explicit evidence boundaries',sources=['docs/README.md','docs/demo_implementation_20260925.md','docs/live_demo_verification_20260927.md'],note='目前已经打通多实例 RGB-D 感知、局部碗沿候选、完整路径预检、固定放置和逐次重新观察，并有白碗绿碗的监督式实机记录。下一步应做候选数量和重新观察策略的对照实验，测量实际抓取与完整任务成功率、耗时、候选尝试数和人工介入次数。同时需要夹持与释放确认、放置区占用检测。当前没有桌面或携带物环境模型，不应推广为未知复杂环境的通用避障。重点是把可运行演示推进为可量化的双碗连续任务。')
card(s,'Demonstrated','Two-instance perception and local geometry\nCandidate-based complete-path planning\nPhysical lifts and placements\nFresh observation after a completed cycle',.55,2.0,5.98,3.37,body_size=21)
card(s,'Next controlled comparison','Single vs. multiple candidates\nFrozen targets vs. fresh observations\nRepeated white / green bowl trials\nPlacement occupancy and grasp sensing',6.8,2.0,5.98,3.37,accent='blue',body_size=21)
text(s,'Measure: physical grasp success · full-cycle success · cycle time · human interventions',.68,5.75,12.0,.57,20,'ink',True)
take(s,'Goal: quantify the benefit of the added functions, not only demonstrate that a run can finish.')

# 17 Appendix
s=slide('Appendix A — Current configuration and interfaces','APPENDIX  /  IMPLEMENTATION REFERENCE','Values describe the current project configuration, not universal operating limits.',sources=['src/rim_locator/config/place_config.json','src/grasp_executor/config/place_config.json','docs/README.md'],note='这一页供导师追问实现时使用。三个包职责保持分离，新增候选分组 JSON 同时携带任务、观测、物体与候选身份。旧单点和 PoseArray 输入仍兼容，但当前完整抓放入口使用 candidate_groups。预抓与抬升 10 厘米，固定放置下降 20 厘米，等待一秒。位姿以 base 或统一后的 base_link 表示，执行工具是 grasp_center。目标有效期上限为 600 秒，但重新观察依然要求真实新帧，不把旧消息修改时间戳后重复使用。')
table(s,['Item','Current setting','Meaning'],[
['Candidate interface','/rim_locator/grasp_candidates','Versioned JSON grouped by object'],
['Candidates per bowl','Up to 8; 4 mm separation','Bounded local contact alternatives'],
['Tool / planning frame','grasp_center / tool0','Connected by fixed TF'],
['Approach / lift','0.10 m','Along base_link Z'],
['Place descent / dwell','0.20 m / 1.0 s','Fixed reference, then release'],
['Target age ceiling','600 s','New capture still required each round']],.55,1.99,[2.65,3.73,5.85],row_h=.63,font_size=17)
take(s,'No added table model or gripper envelope; URDF robot geometry and trajectory checks remain active.')

# 18 Appendix
s=slide('Appendix B — Reproduction and source records','APPENDIX  /  REPRODUCIBILITY','Root-level scripts and matching New_run wrappers use the same implementation.',sources=['New_run/README.md','docs/README.md','docs/live_demo_verification_20260927.md'],note='如需现场复现，先运行硬件准备脚本并保持终端，再用完整抓放的 plan-only 入口检查本次观测，随后才使用 execute 入口。这里提供 9 月 27 日的现场 ROI，移动相机或改变桌面布置后必须重新核对。离线测试入口会构建并运行隔离环境，不连接真机。所有照片保留了原始来源；报告随附逐页来源清单和中文备注，方便回答导师追问。本文引用历史已完成的测试，本次制作 PPT 没有再次运行机械臂。')
label(s,'Offline build and validation',.6,2.02,11)
rect(s,.55,2.35,12.23,.65,'light');text(s,'./run_pipeline_place_offline_test.sh',.78,2.54,11.8,.31,19,font='Liberation Mono')
label(s,'Hardware preparation → preview → execution',.6,3.31,11)
rect(s,.55,3.65,12.23,1.15,'light')
text(s,'./run_robot_prepare.sh\n./run_pipeline_place_test.sh --vision-config <site_roi.json>\n./run_pipeline_place_execute.sh --vision-config <site_roi.json>',.78,3.82,11.8,.87,17,font='Liberation Mono')
text(s,'Site ROI: src/yolo_vision/config/live_pick_roi_20260927.json',.62,5.0,12,.4,16,'muted')
text(s,'Primary evidence: docs/live_demo_verification_20260927.md\nRecorded images and logs: debug_output/live_demo_20260927_gjgip1/',.62,5.67,12,.66,17)
take(s,'English slides · Chinese speaker notes · original experimental images · traceable source files')

RECORDS[0]['title'] = 'RGB-D–Guided Bowl Grasping and Placement with a UR5e'

# Each page includes its Chinese speaking script and exact source paths.
notes=[];manifest=['slide\ttitle\tkind\tpath']
for s,record in zip(P.slides,RECORDS):
    sources=list(dict.fromkeys(record['sources']+record['images']))
    note=f"第 {record['number']} 页：{record['title'] or 'RGB-D–Guided Bowl Grasping and Placement'}\n\n中文讲稿\n{record['note']}\n\n证据来源（相对工作区根目录）\n"+'\n'.join(sources)
    tf=s.notes_slide.notes_text_frame;tf.text=note
    for para in tf.paragraphs:
        para.font.name='Noto Sans CJK SC';para.font.size=Pt(12)
    notes.append(note)
    for src in record['sources']: manifest.append(f"{record['number']}\t{record['title']}\tdocument/data\t{src}")
    for src in dict.fromkeys(record['images']):manifest.append(f"{record['number']}\t{record['title']}\toriginal image\t{src}")

P.save(OUT/(NAME+'.pptx'))
(OUT/'Chinese_Speaker_Notes.txt').write_text(('\n\n'+'='*64+'\n\n').join(notes),encoding='utf-8')
(OUT/'Source_Manifest.tsv').write_text('\n'.join(manifest)+'\n',encoding='utf-8')
(OUT/'slide_content.json').write_text(json.dumps(RECORDS,ensure_ascii=False,indent=2),encoding='utf-8')
assert len(P.slides)==18
print(OUT/(NAME+'.pptx'))
print('18 slides; 18 Chinese notes; original photos only.')
